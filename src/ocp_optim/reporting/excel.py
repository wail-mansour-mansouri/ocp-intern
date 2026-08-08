"""Génération du rapport Excel à cinq feuilles.

Structure imposée par `documentation-ocp/OUTPUT_FORMAT.md` :

  1. Demand Delivery         — analyse par consommateur, avec traçabilité des sources
  2. Phosphoric Production   — production d'acide 29 et d'acide 54
  3. Acid 29 Stocks          — mouvements et bilan matière de l'acide 29
  4. Acid 54 Stocks          — mouvements et bilan matière de l'acide 54
  5. Central Storage Stocks  — opérations sur IR11 et IR12

Le dossier propose ailleurs une variante à sept feuilles, mais elle est archivée
donc obsolète (anomalie A-11, question Q7).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import constants as C
from .. import units
from ..results import Resultat
from ..scenario import Scenario

__all__ = ["FEUILLES_ATTENDUES", "generer_rapport_excel"]

FEUILLES_ATTENDUES: tuple[str, ...] = (
    "Demand Delivery",
    "Phosphoric Production",
    "Acid 29 Stocks",
    "Acid 54 Stocks",
    "Central Storage Stocks",
)

_ENTETE_FOND = PatternFill("solid", fgColor="1F4E79")
_ENTETE_POLICE = Font(bold=True, color="FFFFFF")
_TITRE_POLICE = Font(bold=True, size=13)
_ALERTE_FOND = PatternFill("solid", fgColor="FCE4E4")
_HAUT_GAUCHE = Alignment(vertical="top", wrap_text=True)
_LARGEUR_MAX = 55


def _titre(feuille: Worksheet, ligne: int, texte: str) -> int:
    """Écrit un titre de section et renvoie le numéro de la ligne suivante."""
    cellule = feuille.cell(row=ligne, column=1, value=texte)
    cellule.font = _TITRE_POLICE
    return ligne + 1


def _entete(feuille: Worksheet, ligne: int, colonnes: list[str]) -> int:
    """Écrit une ligne d'en-tête mise en forme."""
    for index, nom in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=ligne, column=index, value=nom)
        cellule.fill = _ENTETE_FOND
        cellule.font = _ENTETE_POLICE
        cellule.alignment = _HAUT_GAUCHE
    return ligne + 1


def _ecrire(feuille: Worksheet, ligne: int, valeurs: list, alerte: bool = False) -> int:
    """Écrit une ligne de données, arrondie au dixième de tonne."""
    for index, valeur in enumerate(valeurs, start=1):
        if isinstance(valeur, float):
            valeur = round(valeur, 1)
        cellule = feuille.cell(row=ligne, column=index, value=valeur)
        cellule.alignment = _HAUT_GAUCHE
        if alerte:
            cellule.fill = _ALERTE_FOND
    return ligne + 1


def _ajuster_largeurs(feuille: Worksheet) -> None:
    """Adapte la largeur de chaque colonne à son contenu, dans une limite raisonnable."""
    for colonne in feuille.columns:
        longueur = max(
            (max(len(l) for l in str(c.value).split("\n")) if c.value is not None else 0)
            for c in colonne
        )
        lettre = get_column_letter(colonne[0].column)
        feuille.column_dimensions[lettre].width = min(longueur + 3, _LARGEUR_MAX)


def _sources_par_acide(resultat: Resultat) -> dict[tuple[str, str], dict[str, float]]:
    """Index inverse : ``(consommateur, acide) -> {ligne source: quantité}``.

    Le format de sortie demande la traçabilité de chaque livraison ; il faut donc
    retourner l'indexation naturelle du modèle, qui est orientée par ligne source.
    """
    index: dict[tuple[str, str], dict[str, float]] = {}

    for ligne, r in resultat.lignes_29.items():
        for consommateur, quantite in r.livraisons_std.items():
            index.setdefault((consommateur, C.ACIDE_29_STD), {})[ligne] = quantite
        for consommateur, quantite in r.livraisons_dec.items():
            index.setdefault((consommateur, C.ACIDE_29_DEC), {})[ligne] = quantite

    for ligne, r in resultat.lignes_54.items():
        for consommateur, quantite in r.livraisons_ncl.items():
            index.setdefault((consommateur, C.ACIDE_54_NCL), {})[ligne] = quantite

    for consommateur, quantite in resultat.ir12.livraisons.items():
        index.setdefault((consommateur, C.ACIDE_54_CL), {})["IR12"] = quantite
    for consommateur, quantite in resultat.ir11.livraisons.items():
        index.setdefault((consommateur, C.BESOIN_54_DEC_TOTAL), {})["IR11"] = quantite

    return index


# ─────────────────────────────────────────────────────────────────────────────
# Feuille 1 — Demand Delivery
# ─────────────────────────────────────────────────────────────────────────────

def _feuille_demande(classeur: Workbook, resultat: Resultat, scenario: Scenario) -> None:
    feuille = classeur.create_sheet("Demand Delivery")
    ligne = _titre(feuille, 1, "ANALYSE DE LA DEMANDE ET DES LIVRAISONS (tonnes de P2O5)")
    ligne = _entete(feuille, ligne + 1, [
        "Ligne", "Produits", "Acide requis", "Acide livré et provenance", "Manque",
    ])

    index = _sources_par_acide(resultat)

    for consommateur in C.CONSOMMATEURS:
        besoins = resultat.besoins.get(consommateur)
        if not besoins:
            continue

        produits = scenario.demande_engrais.get(consommateur, {})
        texte_produits = (
            "\n".join(f"{p} ({q:.0f} t)" for p, q in produits.items())
            if produits else "Consommateur direct"
        )
        texte_requis = "\n".join(f"{a} : {q:.1f} t" for a, q in sorted(besoins.items()))

        lignes_livrees = []
        for acide in sorted(besoins):
            sources = index.get((consommateur, acide), {})
            detail = ", ".join(f"{s}" for s in sorted(sources)) or "—"
            livre = resultat.livraisons[consommateur][acide]
            lignes_livrees.append(f"{acide} : {livre:.1f} t (depuis {detail})")

        manque_total = sum(resultat.manques[consommateur].values())
        texte_manque = f"{manque_total:.1f} t" if manque_total > 1e-4 else "aucun"

        ligne = _ecrire(
            feuille, ligne,
            [consommateur, texte_produits, texte_requis,
             "\n".join(lignes_livrees), texte_manque],
            alerte=manque_total > 1e-4,
        )

    _ajuster_largeurs(feuille)


# ─────────────────────────────────────────────────────────────────────────────
# Feuille 2 — Phosphoric Production
# ─────────────────────────────────────────────────────────────────────────────

def _feuille_production(classeur: Workbook, resultat: Resultat) -> None:
    feuille = classeur.create_sheet("Phosphoric Production")

    ligne = _titre(feuille, 1, "PRODUCTION D'ACIDE 29 % (tonnes de P2O5 par jour)")
    ligne = _entete(feuille, ligne + 1, [
        "Ligne", "Production totale", "Production standard", "Décadmiation",
        "Stock initial STD", "Stock initial DEC",
        "STD vers concentration", "DEC vers concentration",
        "STD vers clients", "DEC vers clients",
        "Transferts entrants", "Transferts sortants",
    ])
    for nom in C.LIGNES_29:
        r = resultat.lignes_29[nom]
        ligne = _ecrire(feuille, ligne, [
            nom, r.production_totale, r.production_std, r.decadmiation,
            r.stock_initial_std, r.stock_initial_dec,
            r.vers_concentration_std, r.vers_concentration_dec,
            r.total_livraisons_std, r.total_livraisons_dec,
            r.transferts_entrants, r.transferts_sortants,
        ])

    ligne += 2
    ligne = _titre(feuille, ligne, "PRODUCTION D'ACIDE 54 % (tonnes de P2O5 par jour)")
    ligne = _entete(feuille, ligne + 1, [
        "Ligne", "Production planifiée", "Charge réelle", "Écart au plan",
        "Entrée cocristallisation", "CoC produit", "NCL produit",
        "Stock initial NCL", "NCL vers clarification", "DEC vers clarification",
        "Décanteurs utilisés (%)", "NCL vers clients",
    ])
    for nom in C.LIGNES_54:
        r = resultat.lignes_54[nom]
        ligne = _ecrire(feuille, ligne, [
            nom, r.production_planifiee, r.charge_reelle,
            r.charge_reelle - r.production_planifiee,
            r.coc_entree, r.coc_produit, r.ncl_produit,
            r.stock_initial_ncl, r.clarification_ncl, r.clarification_dec,
            100.0 * r.taux_occupation_decanteurs, r.total_livraisons_ncl,
        ])

    if resultat.transferts:
        ligne += 2
        ligne = _titre(feuille, ligne, "TRANSFERTS INTERZONES D'ACIDE 29 STANDARD")
        ligne = _entete(feuille, ligne + 1, ["Origine", "Destination", "Quantité"])
        for (origine, destination), quantite in sorted(resultat.transferts.items()):
            ligne = _ecrire(feuille, ligne, [origine, destination, quantite])

    _ajuster_largeurs(feuille)


# ─────────────────────────────────────────────────────────────────────────────
# Feuille 3 — Acid 29 Stocks
# ─────────────────────────────────────────────────────────────────────────────

def _feuille_stocks_29(classeur: Workbook, resultat: Resultat) -> None:
    feuille = classeur.create_sheet("Acid 29 Stocks")
    z_min, z_max = units.bornes_stock_29()

    ligne = _titre(
        feuille, 1,
        f"MOUVEMENTS ET BILAN MATIÈRE DE L'ACIDE 29 % "
        f"(bande de sécurité : {z_min:.1f} – {z_max:.1f} t)"
    )
    ligne = _entete(feuille, ligne + 1, [
        "Ligne", "Type", "Stock initial", "Production", "Transferts entrants",
        "Transferts sortants", "Retours de boue", "Retour EMAPHOS",
        "Vers concentration", "Livraisons", "Stock final", "Écart à la bande",
    ])

    for nom in C.LIGNES_29:
        r = resultat.lignes_29[nom]
        total_final = r.stock_final_std + r.stock_final_dec
        ecart = max(0.0, total_final - z_max) + max(0.0, z_min - total_final)
        alerte = ecart > 0.1

        ligne = _ecrire(feuille, ligne, [
            nom, "acid_29_std", r.stock_initial_std, r.production_std,
            r.transferts_entrants, r.transferts_sortants,
            r.boue_cocristallisation + r.boue_clarification, r.retour_emaphos,
            r.vers_concentration_std, r.total_livraisons_std, r.stock_final_std,
            "" if not alerte else round(ecart, 1),
        ], alerte=alerte)
        ligne = _ecrire(feuille, ligne, [
            nom, "acid_29_dec", r.stock_initial_dec, r.decadmiation,
            0.0, 0.0, 0.0, 0.0,
            r.vers_concentration_dec, r.total_livraisons_dec, r.stock_final_dec, "",
        ], alerte=alerte)

    ligne += 1
    ligne = _ecrire(feuille, ligne, [
        "Note", "La bande de sécurité porte sur la SOMME des deux qualités : "
        "elles partagent les mêmes cuves.",
    ])
    _ajuster_largeurs(feuille)


# ─────────────────────────────────────────────────────────────────────────────
# Feuille 4 — Acid 54 Stocks
# ─────────────────────────────────────────────────────────────────────────────

def _feuille_stocks_54(classeur: Workbook, resultat: Resultat) -> None:
    feuille = classeur.create_sheet("Acid 54 Stocks")
    z_min, z_max = units.bornes_stock_54()

    ligne = _titre(
        feuille, 1,
        f"MOUVEMENTS ET BILAN MATIÈRE DE L'ACIDE 54 % NCL "
        f"(bande de sécurité : {z_min:.1f} – {z_max:.1f} t)"
    )
    ligne = _entete(feuille, ligne + 1, [
        "Ligne", "Stock initial", "NCL produit", "Vers cocristallisation",
        "Vers clarification NCL", "Vers clarification DEC",
        "Livraisons", "Stock final", "Écart à la bande",
    ])

    for nom in C.LIGNES_54:
        r = resultat.lignes_54[nom]
        ecart = max(0.0, r.stock_final_ncl - z_max) + max(0.0, z_min - r.stock_final_ncl)
        alerte = ecart > 0.1
        ligne = _ecrire(feuille, ligne, [
            nom, r.stock_initial_ncl, r.ncl_produit, r.coc_entree,
            r.clarification_ncl, r.clarification_dec,
            r.total_livraisons_ncl, r.stock_final_ncl,
            "" if not alerte else round(ecart, 1),
        ], alerte=alerte)

    ligne += 1
    ligne = _ecrire(feuille, ligne, [
        "Note", "Aucune boue ne revient à ce niveau : elles remontent au stock "
        "d'acide 29. Le NCL décadmié n'est jamais stocké localement.",
    ])
    _ajuster_largeurs(feuille)


# ─────────────────────────────────────────────────────────────────────────────
# Feuille 5 — Central Storage Stocks
# ─────────────────────────────────────────────────────────────────────────────

def _feuille_stockage_central(classeur: Workbook, resultat: Resultat) -> None:
    feuille = classeur.create_sheet("Central Storage Stocks")
    ligne = _titre(feuille, 1, "STOCKAGE CENTRAL — IR11 ET IR12 (tonnes de P2O5)")
    ligne = _entete(feuille, ligne + 1, [
        "Bac", "Contenu", "Stock initial", "Entrée CoC", "Entrée DEC_CL",
        "Entrée CL", "Total entrées", "Livraisons", "Stock final",
        "Capacité", "Remplissage (%)", "Dépassement",
    ])

    for bac, contenu in ((resultat.ir11, "CoC + DEC_CL"), (resultat.ir12, "CL")):
        depassement = max(0.0, bac.stock_final - bac.capacite)
        ligne = _ecrire(feuille, ligne, [
            bac.nom, contenu, bac.stock_initial,
            bac.entrees.get("CoC", 0.0), bac.entrees.get("DEC_CL", 0.0),
            bac.entrees.get("CL", 0.0), bac.total_entrees,
            bac.total_livraisons, bac.stock_final, bac.capacite,
            100.0 * bac.taux_remplissage,
            "" if depassement <= 0.1 else round(depassement, 1),
        ], alerte=depassement > 0.1)

    ligne += 2
    ligne = _titre(feuille, ligne, "LIVRAISONS DEPUIS LE STOCKAGE CENTRAL")
    ligne = _entete(feuille, ligne + 1, ["Bac", "Consommateur", "Quantité"])
    for bac in (resultat.ir11, resultat.ir12):
        for consommateur, quantite in sorted(bac.livraisons.items()):
            ligne = _ecrire(feuille, ligne, [bac.nom, consommateur, quantite])

    _ajuster_largeurs(feuille)


# ─────────────────────────────────────────────────────────────────────────────
# Assemblage
# ─────────────────────────────────────────────────────────────────────────────

def generer_rapport_excel(
    resultat: Resultat, scenario: Scenario, chemin: Path | str
) -> Path:
    """Écrit le classeur Excel à cinq feuilles.

    Args:
        resultat: solution structurée.
        scenario: scénario d'origine (pour les libellés de produits).
        chemin: fichier de sortie (`.xlsx`).

    Returns:
        Le chemin du fichier écrit.
    """
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)

    classeur = Workbook()
    classeur.remove(classeur.active)   # retire la feuille vide créée par défaut

    _feuille_demande(classeur, resultat, scenario)
    _feuille_production(classeur, resultat)
    _feuille_stocks_29(classeur, resultat)
    _feuille_stocks_54(classeur, resultat)
    _feuille_stockage_central(classeur, resultat)

    classeur.save(chemin)
    return chemin
